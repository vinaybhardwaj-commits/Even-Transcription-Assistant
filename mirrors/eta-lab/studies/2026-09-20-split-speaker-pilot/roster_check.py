#!/usr/bin/env python3
"""Cross-check every claim against the OPD roster. No listening, no audio, no clip content.

A ROSTER IS A PLAN, NOT AN ATTENDANCE RECORD. Doctors swap, run late, finish early and cover for
each other, so a claim outside a slot is UNEXPLAINED, never "wrong". The buckets are named
consistent / unexplained / ambiguous for that reason, and a large unexplained bucket is as much
evidence that the roster is departed from as that the matcher erred.
"""
import datetime, json, re, sys, collections
import openpyxl
from q import q

XLSX = '/Users/vinaybhardwaj/Desktop/SEPT 2026 OPD Schedule.xlsx'
OUT = '/Users/vinaybhardwaj/dev/Even-Transcription-Assistant/docs/handoff/scratch/roster-crosscheck-20-SEP-2026.json'
IST = datetime.timezone(datetime.timedelta(hours=5, minutes=30))
DAYS = ['MONDAY','TUESDAY','WEDNESDAY','THURSDAY','FRIDAY','SATURDAY','SUNDAY']

def parse_time(t):
    """'3:30PM', '11AM' -> minutes since midnight. None if unparseable."""
    m = re.match(r'^\s*(\d{1,2})(?::(\d{2}))?\s*(AM|PM)\s*$', t, re.I)
    if not m: return None
    h, mi, ap = int(m.group(1)), int(m.group(2) or 0), m.group(3).upper()
    if h == 12: h = 0
    return (h + (12 if ap == 'PM' else 0)) * 60 + mi

def parse_range(s):
    if not s: return None
    parts = re.split(r'\s*[-–]\s*', str(s).strip())
    if len(parts) != 2: return None
    a, b = parse_time(parts[0]), parse_time(parts[1])
    return None if a is None or b is None else (a, b)

def parse_docs(s):
    """'Dr Animesh/ Dr Rakshitha (ENT) - Walk in' -> ['Animesh','Rakshitha']; FREE -> []."""
    if not s: return []
    s = str(s).strip()
    if re.match(r'^\s*FREE\s*$', s, re.I): return []
    s = re.sub(r'\([^)]*\)', ' ', s)          # drop specialty
    s = re.split(r'\s+-\s+', s)[0]            # drop '- Walk in' / '- Remote TC'
    out = []
    for piece in s.split('/'):
        p = re.sub(r'\bDr\.?\b', ' ', piece, flags=re.I).strip(' .,')
        p = re.sub(r'\s+', ' ', p)
        if p: out.append(p)
    return out

def load_roster():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    grid = [[c.value for c in row] for row in ws.iter_rows()]
    # room header columns from row 1
    rooms = {}
    for j, v in enumerate(grid[0]):
        if v and re.match(r'^\s*ROOM\s*\d+', str(v), re.I):
            rooms[j] = re.sub(r'\s+', ' ', str(v).strip().upper())
    # weekday blocks from column A
    starts = [(i, str(grid[i][0]).strip().upper()) for i in range(len(grid))
              if grid[i] and grid[i][0] and str(grid[i][0]).strip().upper() in DAYS]
    slots, unparsed = [], []
    for bi, (r0, day) in enumerate(starts):
        r1 = starts[bi+1][0] if bi+1 < len(starts) else len(grid)
        for j, rname in rooms.items():
            for r in range(r0, r1):
                trng = grid[r][j] if j < len(grid[r]) else None
                dcell = grid[r][j+1] if j+1 < len(grid[r]) else None
                if not trng and not dcell: continue
                rng = parse_range(trng)
                if rng is None:
                    if str(trng or '').strip(): unparsed.append(dict(day=day, room=rname, cell=str(trng)[:40]))
                    continue
                docs = parse_docs(dcell)
                slots.append(dict(day=day, room=rname, start_min=rng[0], end_min=rng[1],
                                  doctors=docs, free=(not docs)))
    return rooms, slots, unparsed

def rows(sql, p=None):
    r = q(sql, p); return r['rows'] if isinstance(r, dict) and 'rows' in r else r

def main():
    rooms, slots, unparsed = load_roster()
    clin = rows("select id, full_name, given_name, family_name from clinician")
    vp = {x['doctor_id'] for x in rows("select distinct doctor_id from voice_print")}
    # roster name -> clinician_id. CONSERVATIVE: exact normalised match, or the roster's tokens a
    # subset of the record's tokens. Anything ambiguous or near-miss stays UNMATCHED and is reported.
    # A wrong match here would attribute one doctor's audio to another, which is unrecoverable.
    def norm(x):
        x = re.sub(r'\bdrs?\b', ' ', str(x or '').lower())
        x = re.sub(r'[^a-z ]', ' ', x)
        return re.sub(r'\s+', ' ', x).strip()
    def match(name):
        n = norm(name); nt = set(n.split())
        if not nt: return None
        exact = [c for c in clin if norm(c['full_name']) == n]
        if len(exact) == 1: return exact[0]['id']
        sub = [c for c in clin if nt and nt <= set(norm(c['full_name']).split())]
        return sub[0]['id'] if len(sub) == 1 else None
    rostered = sorted({d for s in slots for d in s['doctors']})
    name2id = {d: match(d) for d in rostered}

    # DB rooms -> schedule room, proven from room.name
    droom = rows("select id, name from room")
    room_map, room_excl = {}, []
    for r in droom:
        m = re.match(r'^\s*OPD\s*(\d+)', str(r['name'] or ''), re.I)
        if m and ('ROOM %s' % m.group(1)) in rooms.values():
            room_map[r['id']] = 'ROOM %s' % m.group(1)
        else:
            room_excl.append(dict(id=r['id'], name=r['name']))

    # every claim
    claims = rows("""select t.clinician_id, rd.room_id, rd.ist_date::text d,
                            split_part(t.source_ref,'|',2)::bigint s
                     from room_turn_speaker t join room_day rd on rd.id=t.room_day_id
                     where t.clinician_id is not null""")
    unnamed = rows("""select rd.room_id, rd.ist_date::text d,
                             split_part(t.source_ref,'|',2)::bigint s
                      from room_turn_speaker t join room_day rd on rd.id=t.room_day_id
                      where t.clinician_id is null""")
    id2name = {v: k for k, v in name2id.items() if v}

    def slots_for(room_id, ms):
        dt = datetime.datetime.fromtimestamp(int(ms)/1000, IST)
        day = DAYS[dt.weekday()]; mins = dt.hour*60 + dt.minute
        rn = room_map.get(room_id)
        if rn is None: return None, None
        return [s for s in slots if s['day']==day and s['room']==rn and s['start_min']<=mins<s['end_min']], rn

    per = collections.defaultdict(lambda: collections.Counter())
    for c in claims:
        sl, rn = slots_for(c['room_id'], c['s'])
        key = (c['clinician_id'], rn)
        if sl is None: per[key]['room_unmapped'] += 1; continue
        per[key]['n'] += 1
        mine = [s for s in sl if any(name2id.get(d) == c['clinician_id'] for d in s['doctors'])]
        if not sl: per[key]['unexplained_noslot'] += 1
        elif mine and len(mine[0]['doctors']) > 1: per[key]['ambiguous_shared'] += 1
        elif mine: per[key]['consistent'] += 1
        else: per[key]['unexplained_notrostered'] += 1

    miss = collections.Counter()
    for u in unnamed:
        sl, rn = slots_for(u['room_id'], u['s'])
        if sl is None or not sl: continue
        if len(sl[0]['doctors']) == 1 and name2id.get(sl[0]['doctors'][0]):
            miss[(name2id[sl[0]['doctors'][0]], rn)] += 1

    json.dump(dict(rooms=rooms, slots=slots, unparsed=unparsed, rostered=rostered,
                   name2id=name2id, room_map=room_map, room_excluded=room_excl,
                   per={'%s|%s'%k: dict(v) for k,v in per.items()},
                   misses={'%s|%s'%k: v for k,v in miss.items()},
                   voiceprint_ids=sorted(vp)), open(OUT,'w'), indent=1, default=str)

    print('ROSTER: %d slots, %d rooms, %d weekday blocks; unparsed time cells: %d'
          % (len(slots), len(rooms), len({s['day'] for s in slots}), len(unparsed)))
    print('rostered names: %d | matched to clinician_id: %d | with voiceprint: %d'
          % (len(rostered), sum(1 for v in name2id.values() if v),
             sum(1 for v in name2id.values() if v in vp)))
    print('\nPER CLINICIAN x ROOM (claims)')
    print('  %-14s %-7s %5s %6s %6s %6s %6s' % ('clinician','room','n','consis','shared','no-slot','not-ros'))
    for (cid, rn), v in sorted(per.items(), key=lambda kv: -kv[1]['n']):
        print('  %-14s %-7s %5d %6d %6d %6d %6d' % (cid, rn or '?', v['n'], v['consistent'],
              v['ambiguous_shared'], v['unexplained_noslot'], v['unexplained_notrostered']))
    print('\nMISSES (unnamed turns inside a solo-rostered slot)')
    for (cid, rn), n in sorted(miss.items(), key=lambda kv:-kv[1]): print('  %-14s %-7s %5d' % (cid, rn, n))
    print('\nEXCLUSIONS')
    print('  rooms not mapped to a schedule column: %d' % len(room_excl))
    for r in room_excl: print('    %s  %s' % (r['id'], r['name']))
    print('  rostered names with NO voiceprint: %s'
          % ', '.join(sorted(d for d in rostered if name2id.get(d) not in vp)) or 'none')
    print('  enrolled voiceprints not on the roster: %s'
          % ', '.join(sorted(i for i in vp if i not in id2name)) or 'none')
    for u in unparsed[:6]: print('  unparsed cell: %s %s %r' % (u['day'], u['room'], u['cell']))
    print('\ndetail -> %s' % OUT)

main()
